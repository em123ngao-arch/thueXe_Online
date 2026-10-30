"""
Script sinh file trình chiếu PowerPoint (19 Slides tỷ lệ 16:9)
Bảo vệ đồ án tốt nghiệp Khóa K25 - Đề tài DriveShare
- Thiết kế chuẩn mẫu bài thuyết trình ITC (docs/img_slide):
  + Biểu tượng '>>>' màu cam nổi bật ở tiêu đề
  + Đơn giản, cực kỳ ít chữ, hình ảnh & sơ đồ là trung tâm
  + Thể hiện 4 bảng CSDL cốt lõi (users, cars, rentals, payments)
  + 3 slide phân hệ (Khách thuê, Chủ xe, Admin) đều có 3 ảnh giao diện nằm ngang
  + Nhúng các sơ đồ Use Case, State Machine, Sequence Diagram to rõ ràng
  + 100% tiếng Việt có dấu, thông tin chuẩn nhóm K25
  + Speaker Notes đầy đủ 15 phút từng giây
"""

import os
import sys

# Thiết lập UTF-8 cho console Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Bảng màu chuẩn Doanh nghiệp học thuật ITC
PRIMARY = RGBColor(30, 58, 138)        # Xanh Navy đậm (#1E3A8A)
PRIMARY_DARK = RGBColor(15, 23, 42)    # Xanh đen Slate (#0F172A)
PRIMARY_LIGHT = RGBColor(238, 242, 255)# Xanh pastel nhạt (#EEF2FF)
ACCENT_ORANGE = RGBColor(234, 88, 12)  # Cam điểm nhấn ITC (>>>) (#EA580C)
ACCENT_TEAL = RGBColor(13, 148, 136)   # Xanh mòng két (#0D9488)
BG_LIGHT = RGBColor(248, 250, 252)     # Nền xám sáng tinh khôi (#F8FAFC)
WHITE = RGBColor(255, 255, 255)        # Trắng thuần
TEXT_MAIN = RGBColor(15, 23, 42)       # Chữ đậm (#0F172A)
TEXT_MUTED = RGBColor(71, 85, 105)     # Chữ trung tính (#475569)
BORDER_COLOR = RGBColor(226, 232, 240) # Viền xám nhạt (#E2E8F0)
CODE_BG = RGBColor(15, 23, 42)         # Nền code tối (#0F172A)
CODE_YELLOW = RGBColor(250, 204, 21)   # Vàng highlight (#FACC15)
CODE_GREEN = RGBColor(74, 222, 128)    # Xanh lá highlight (#4ADE80)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"


def set_slide_background(slide, color=BG_LIGHT):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_slide_header(slide, section_number, title_text, subtitle_text=None):
    """
    Header chuẩn theo phong cách bài mẫu ITC (img_slide):
    - Dải cam mảnh ở mép trên
    - Biểu tượng >>> màu cam + Số mục + Tiêu đề lớn
    - Phụ đề nhỏ dẫn dắt (nếu có)
    """
    # Top accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_ORANGE
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    # '>>>' cam
    run_icon = p.add_run()
    run_icon.text = ">>>  "
    run_icon.font.name = FONT_HEADING
    run_icon.font.size = Pt(22)
    run_icon.font.bold = True
    run_icon.font.color.rgb = ACCENT_ORANGE

    # Tên mục
    run_title = p.add_run()
    run_title.text = f"{section_number}  {title_text}"
    run_title.font.name = FONT_HEADING
    run_title.font.size = Pt(23)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_DARK

    if subtitle_text:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_before = Pt(4)


def set_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text.strip()


def add_centered_diagram_slide(slide, section_no, title, subtitle, img_path, caption):
    """
    Slide hiển thị 1 sơ đồ lớn ở trung tâm (Use Case, State Machine, Sequence Diagram, Bảng CSDL)
    Chuẩn theo phong cách bài mẫu của ITC: Rất ít chữ, sơ đồ chiếm 80% diện tích.
    """
    set_slide_background(slide)
    add_slide_header(slide, section_no, title, subtitle)

    box_l = Inches(0.8)
    box_t = Inches(1.65)
    box_w = Inches(11.733)
    box_h = Inches(4.95)

    # Khung nền trắng bo góc nhẹ cho ảnh
    container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_l, box_t, box_w, box_h)
    container.fill.solid()
    container.fill.fore_color.rgb = WHITE
    container.line.color.rgb = BORDER_COLOR
    container.line.width = Pt(1.5)

    if img_path and os.path.exists(img_path):
        try:
            with Image.open(img_path) as img:
                orig_w, orig_h = img.size

            aspect = orig_w / orig_h
            avail_w = box_w - Inches(0.4)
            avail_h = box_h - Inches(0.4)
            container_aspect = avail_w / avail_h

            if aspect > container_aspect:
                dest_w = avail_w
                dest_h = dest_w / aspect
            else:
                dest_h = avail_h
                dest_w = dest_h * aspect

            dest_l = box_l + (box_w - dest_w) / 2
            dest_t = box_t + (box_h - dest_h) / 2

            slide.shapes.add_picture(img_path, dest_l, dest_t, dest_w, dest_h)
        except Exception as e:
            print(f"Lỗi chèn ảnh {img_path}: {e}")

    # Dòng caption phía dưới
    tb_cap = slide.shapes.add_textbox(box_l, Inches(6.75), box_w, Inches(0.4))
    tf_c = tb_cap.text_frame
    p_c = tf_c.paragraphs[0]
    p_c.alignment = PP_ALIGN.CENTER
    p_c.text = caption
    p_c.font.name = FONT_BODY
    p_c.font.size = Pt(12)
    p_c.font.italic = True
    p_c.font.color.rgb = TEXT_MUTED


def add_three_mobile_screens_slide(slide, section_no, title, subtitle, screens_data):
    """
    Slide hiển thị 3 ảnh màn hình giao diện (Mobile/Web) đặt nằm ngang cạnh nhau
    screens_data = [ (img_path, caption_title, caption_desc), ... ]
    """
    set_slide_background(slide)
    add_slide_header(slide, section_no, title, subtitle)

    col_w = Inches(3.7)
    gap = Inches(0.316)
    box_t = Inches(1.65)
    box_h = Inches(4.75)

    for i, (img_path, cap_title, cap_desc) in enumerate(screens_data):
        c_left = Inches(0.8 + i * (3.7 + 0.316))

        # Card container trắng
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, box_t, col_w, box_h)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        # Chèn ảnh bên trong card
        if img_path and os.path.exists(img_path):
            try:
                with Image.open(img_path) as img:
                    orig_w, orig_h = img.size

                aspect = orig_w / orig_h
                avail_w = col_w - Inches(0.24)
                avail_h = box_h - Inches(0.9)  # chừa chỗ cho caption
                container_aspect = avail_w / avail_h

                if aspect > container_aspect:
                    dest_w = avail_w
                    dest_h = dest_w / aspect
                else:
                    dest_h = avail_h
                    dest_w = dest_h * aspect

                dest_l = c_left + (col_w - dest_w) / 2
                dest_t = box_t + Inches(0.12) + (avail_h - dest_h) / 2

                slide.shapes.add_picture(img_path, dest_l, dest_t, dest_w, dest_h)
            except Exception as e:
                print(f"Lỗi chèn ảnh {img_path}: {e}")

        # Caption phía dưới ảnh
        tb_cap = slide.shapes.add_textbox(c_left + Inches(0.1), box_t + box_h - Inches(0.85), col_w - Inches(0.2), Inches(0.75))
        tf = tb_cap.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = cap_title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = cap_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED


def add_three_cards_slide(slide, section_no, title, subtitle, cards_data):
    """Slide 3 Thẻ Card bo góc nằm ngang súc tích"""
    set_slide_background(slide)
    add_slide_header(slide, section_no, title, subtitle)

    for i, (tag_card, ctitle, cbullets) in enumerate(cards_data):
        left_pos = Inches(0.8 + i * 4.0)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.75), Inches(3.733), Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.35)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag_card
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_ORANGE
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = ctitle
        p.font.name = FONT_HEADING
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(14)

        for b in cbullets:
            p = tf.add_paragraph()
            p.text = "•  " + b
            p.font.name = FONT_BODY
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(9)


def add_four_cards_slide(slide, section_no, title, subtitle, cards_data):
    """Slide Lưới 2x2: Chữ to rõ, thoáng đãng"""
    set_slide_background(slide)
    add_slide_header(slide, section_no, title, subtitle)

    coords = [
        (Inches(0.8), Inches(1.75)),
        (Inches(6.8), Inches(1.75)),
        (Inches(0.8), Inches(4.45)),
        (Inches(6.8), Inches(4.45)),
    ]

    for i, (tag_card, ctitle, cdesc) in enumerate(cards_data):
        c_left, c_top = coords[i]
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(5.733), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag_card
        p.font.name = FONT_HEADING
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_ORANGE
        p.space_after = Pt(2)

        p = tf.add_paragraph()
        p.text = ctitle
        p.font.name = FONT_HEADING
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(5)

        lines = cdesc.split("\n")
        for idx, line in enumerate(lines):
            p = tf.add_paragraph()
            p.text = line
            p.font.name = FONT_BODY
            p.font.size = Pt(12.5)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(2)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TRANG BÌA (KHÓA K25, GVHD & 6 THÀNH VIÊN THEO ĐÚNG USER ĐÃ SỬA)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.18))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT_ORANGE
    top_bar.line.fill.background()

    tb_tag = s1.shapes.add_textbox(Inches(1.0), Inches(0.75), Inches(11.3), Inches(0.4))
    p = tb_tag.text_frame.paragraphs[0]
    p.text = "KHOA CÔNG NGHỆ THÔNG TIN - ITC  |  KHÓA K25"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    tb_main = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(2.2))
    tf = tb_main.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "DRIVESHARE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    p2 = tf.add_paragraph()
    p2.text = "Nền Tảng Cho Thuê Xe Tự Lái P2P Trực Tuyến"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_TEAL

    p3 = tf.add_paragraph()
    p3.text = "Spring Boot 3  •  PostgreSQL 16  •  VietQR Escrow 30%  •  Gemini AI Assistant"
    p3.font.name = FONT_HEADING
    p3.font.size = Pt(14.5)
    p3.font.color.rgb = TEXT_MUTED

    # Card Giảng viên hướng dẫn
    card_gvhd = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.75), Inches(5.3), Inches(3.15))
    card_gvhd.fill.solid()
    card_gvhd.fill.fore_color.rgb = WHITE
    card_gvhd.line.color.rgb = BORDER_COLOR
    card_gvhd.line.width = Pt(1.5)

    tf_gv = card_gvhd.text_frame
    tf_gv.margin_left = tf_gv.margin_right = tf_gv.margin_top = Inches(0.35)
    p = tf_gv.paragraphs[0]
    p.text = "GIẢNG VIÊN HƯỚNG DẪN"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(4)

    p = tf_gv.add_paragraph()
    p.text = "Thầy Trương Thanh Bình"
    p.font.name = FONT_HEADING
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_after = Pt(12)

    p = tf_gv.add_paragraph()
    p.text = "Bộ môn: Công nghệ Phần mềm\nTrường Cao đẳng Công nghệ Thông tin TP.HCM (ITC)\nNiên khóa: 2025 - 2027"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED

    # Card Nhóm sinh viên thực hiện (K25 - 6 thành viên chuẩn)
    card_team = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(3.75), Inches(5.7), Inches(3.15))
    card_team.fill.solid()
    card_team.fill.fore_color.rgb = WHITE
    card_team.line.color.rgb = BORDER_COLOR
    card_team.line.width = Pt(1.5)

    tf_team = card_team.text_frame
    tf_team.margin_left = tf_team.margin_right = tf_team.margin_top = Inches(0.35)
    p = tf_team.paragraphs[0]
    p.text = "NHÓM THỰC HIỆN (6 THÀNH VIÊN)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(6)

    members = [
        ("Nguyễn Duy Bảo", "Tech Lead / Thuyết trình chính"),
        ("Nguyễn Chí Tín", "Fullstack Dev (Payment & UI)"),
        ("Lâm Chí Vĩ", "Backend Core (Handover & Locking)"),
        ("Đỗ Ngọc Phát", "AI & Security (Gemini & Token)"),
        ("Nguyễn Duy Quân", "QA Tester (188 Tests & Overdue)"),
        ("NguyễnTấn Lộc Khiêm", "Kiến trúc sư (Use Case & ERD)")
    ]
    for name, role in members:
        p = tf_team.add_paragraph()
        p.text = f"•  {name}  —  {role}"
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN if "Bảo" in name else TEXT_MUTED

    set_notes(s1, """LỜI THOẠI THUYẾT TRÌNH (00:00 - 00:30):
Kính thưa quý Thầy Cô trong Hội đồng phản biện và Thầy Trương Thanh Bình kính mến!
Em tên là Nguyễn Duy Bảo, đại diện cho nhóm sinh viên Khóa K25 chuyên ngành Công nghệ Phần mềm.
Hôm nay, nhóm chúng em xin được báo cáo đề tài tốt nghiệp: DriveShare - Nền tảng chia sẻ và cho thuê xe tự lái trực tuyến theo mô hình P2P.
Bài thuyết trình gồm 3 chương trọng tâm với thời lượng 15 phút. Em xin phép được bắt đầu!""")

    # =========================================================================
    # SLIDE 2: CHƯƠNG 1 - BỐI CẢNH & BẤT CẬP THUÊ TRUYỀN THỐNG (3 CARDS)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    card_data_s2 = [
        ("BẤT CẬP 01", "Rủi Ro Giữ Tài Sản",
         [
             "Bắt buộc giữ CCCD gốc hoặc xe máy kèm cà-vẹt đắt tiền.",
             "Nguy cơ lộ lọt thông tin cá nhân và thiếu căn cứ pháp lý.",
             "Gây áp lực tâm lý và rào cản tài chính lớn cho khách thuê."
         ]),
        ("BẤT CẬP 02", "Tranh Chấp Hư Hỏng",
         [
             "Không có quy trình nghiệm thu tình trạng xe minh bạch.",
             "Khách thuê dễ bị 'bắt đền oan' các vết xước có từ trước.",
             "Thiếu bằng chứng đối soát khi hai bên bàn giao chìa khóa."
         ]),
        ("BẤT CẬP 03", "Quản Lý Thủ Công",
         [
             "Chốt đơn qua tin nhắn Zalo, sổ tay nên rất dễ quên lịch.",
             "Nguy cơ nhận cọc của 2 người cùng một ngày (Overbooking).",
             "Chưa có công cụ tự động hóa kiểm soát xe quá hạn hợp đồng."
         ])
    ]
    add_three_cards_slide(s2, "1.1", "Bối Cảnh & Một Vài Bất Cập Của Thuê Xe Truyền Thống",
                          "> Nhu cầu thuê xe tự lái gia tăng nhanh nhưng phương thức truyền thống lộ rõ rủi ro:",
                          card_data_s2)
    set_notes(s2, """LỜI THOẠI THUYẾT TRÌNH (00:30 - 01:15):
Thưa Thầy Cô, nhu cầu thuê xe tự lái tăng rất nhanh, nhưng cách thuê truyền thống lộ rõ 3 bất cập lớn:
1. Rủi ro giam giữ tài sản: Khách phải để lại CCCD gốc hoặc xe máy giá trị cao.
2. Tranh chấp trầy xước: Thiếu biên bản nghiệm thu nên hay xảy ra cãi vã khi trả xe.
3. Quản lý thủ công: Nhận cọc qua Zalo dễ dẫn đến đặt trùng lịch 2 người cùng ngày.
Mô hình P2P của DriveShare ra đời nhằm giải quyết triệt để các vấn đề này.""")

    # =========================================================================
    # SLIDE 3: CHƯƠNG 1 - BẢNG ĐỐI CHUẨN TIÊU CHÍ (BENCHMARK TABLE)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_slide_header(s3, "1.2", "Bảng Đối Chuẩn Tiêu Chí: Mioto vs Zoomcar vs DriveShare",
                     "> So sánh 5 tiêu chí nghiệp vụ cốt lõi giữa các nền tảng cho thuê xe:")

    rows = 6
    cols = 4
    left = Inches(0.8)
    top = Inches(1.75)
    width = Inches(11.733)
    height = Inches(5.1)

    table_shape = s3.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(2.7)
    table.columns[2].width = Inches(2.7)
    table.columns[3].width = Inches(3.133)

    headers = ["Tiêu Chí Nghiệp Vụ", "Mioto", "Zoomcar", "DriveShare (Đề Tài K25)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            p.font.name = FONT_HEADING
            p.font.size = Pt(13.5)
            p.font.bold = True
            p.font.color.rgb = WHITE

    benchmark_rows = [
        ("Chính sách Cọc & Giữ chỗ", "100% tiền thuê chuyển khoản", "Giữ qua Thẻ tín dụng", "Ký quỹ 30% VietQR linh hoạt"),
        ("Trợ lý ảo AI tư vấn 24/7", "Không có", "Không có", "AI Gemini 2.5 Flash + Heuristic"),
        ("Chặn lịch bận xe (Blackout)", "Có hỗ trợ", "Hạn chế tính năng", "Chặn lịch bảo dưỡng chủ động"),
        ("Biên bản bàn giao 2 chiều", "Giấy tờ hoặc 1 bên duyệt", "Nhân viên kiểm tra", "Điện tử 2 chiều: ODO, Xăng, 6 ảnh"),
        ("Chế tài quá hạn & Điều 175", "Tính phí giờ thủ công", "Phạt theo giờ cố định", "Tự động phạt 1.5x-2x & Hồ sơ 1-click")
    ]

    for i, row in enumerate(benchmark_rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else PRIMARY_LIGHT
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                p.font.name = FONT_BODY
                p.font.size = Pt(13)
                p.font.color.rgb = PRIMARY_DARK if j == 3 else TEXT_MAIN
                if j == 3:
                    p.font.bold = True

    set_notes(s3, """LỜI THOẠI THUYẾT TRÌNH (01:15 - 02:15):
Để tạo ưu thế cạnh tranh, nhóm đã lập Bảng đối chuẩn giữa Mioto, Zoomcar với DriveShare qua 5 tiêu chí:
- Về cọc: DriveShare áp dụng Ký quỹ 30% VietQR giúp giảm áp lực tiền mặt.
- Về hỗ trợ: Có Trợ lý AI tính chi phí và tìm xe tự động.
- Về bàn giao: Bắt buộc Check-in/Check-out 2 chiều có ảnh đối soát ODO và xăng.
- Đặc biệt: Tiên phong tự động hóa xử phạt quá hạn và trích xuất hồ sơ vi phạm theo Điều 175 BLHS.""")

    # =========================================================================
    # SLIDE 4: CHƯƠNG 1 - 4 MỤC TIÊU CỐT LÕI (LƯỚI 2X2 ĐƠN GIẢN)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    card_data_s4 = [
        ("MỤC TIÊU 01", "Nền Tảng P2P An Toàn & Tiện Lợi",
         "Xây dựng Web Fullstack Responsive kết nối trực tiếp xe nhàn rỗi và khách thuê với chi phí tối ưu nhất."),
        ("MỤC TIÊU 02", "Bảo Vệ Tài Chính (Escrow 30%)",
         "Sàn đóng vai trò trung gian giữ cọc 30% qua VietQR, chỉ giải ngân cho chủ xe khi cuốc xe hoàn tất."),
        ("MỤC TIÊU 03", "Triệt Tiêu Xung Đột Trùng Lịch",
         "Ứng dụng Khóa bi quan CSDL và Khóa 2 tầng (Two-Tier Lock) loại bỏ hoàn toàn nguy cơ đặt trùng xe."),
        ("MỤC TIÊU 04", "Trợ Lý Ảo AI Thông Minh",
         "Tích hợp mô hình AI Google Gemini 2.5 Flash hỗ trợ khách tìm xe, tính giá cọc và giải đáp 24/7.")
    ]
    add_four_cards_slide(s4, "1.3", "4 Mục Tiêu Cốt Lõi Của Đề Tài",
                         "> Định hướng nghiên cứu và phát triển giải pháp toàn diện cho nền tảng:",
                         card_data_s4)
    set_notes(s4, """LỜI THOẠI THUYẾT TRÌNH (02:15 - 03:00):
Từ các phân tích trên, DriveShare xác định 4 mục tiêu then chốt:
1. Xây dựng nền tảng P2P an toàn trên cả PC và thiết bị di động.
2. Bảo vệ dòng tiền cho cả hai bên nhờ cơ chế giữ cọc Escrow 30%.
3. Giải quyết triệt để lỗi trùng xe bằng khóa CSDL chuyên sâu.
4. Tích hợp AI nâng cao trải nghiệm đặt xe thông minh.""")

    # =========================================================================
    # SLIDE 5: CHƯƠNG 2 - MÔ HÌNH KIẾN TRÚC TỔNG QUAN & TECH STACK
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_slide_header(s5, "2.1", "Mô Hình Kiến Trúc Hệ Thống & Tech Stack",
                     "> Kiến trúc phân tầng Decoupled Client - Server tách biệt qua RESTful API:")

    # Cột trái: Card đặc điểm kiến trúc
    left_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.8), Inches(5.1))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = BORDER_COLOR
    left_card.line.width = Pt(1.5)

    tf_l = left_card.text_frame
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.35)
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "ĐẶC ĐIỂM KIẾN TRÚC CỐT LÕI"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    bullets_s5 = [
        "Kiến trúc Decoupled: Tách biệt Client - Server hoàn toàn qua chuẩn REST API.",
        "Backend Core: Java 17/21, Spring Boot 3.3.4 Layered Modular (9 modules độc lập).",
        "Cơ sở dữ liệu: PostgreSQL 16 Alpine đóng gói Docker (Port 5432) + Adminer.",
        "Lưu trữ Hybrid: Ưu tiên Cloudinary CDN; tự động dự phòng ổ cứng cục bộ.",
        "Trí tuệ nhân tạo: Google Gemini 2.5 Flash REST API + Chat Memory PostgreSQL.",
        "Frontend Web: HTML5, CSS3, JavaScript ES6+, Bootstrap 5 tối ưu tải tĩnh."
    ]
    for b in bullets_s5:
        p = tf_l.add_paragraph()
        p.text = "•  " + b
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(9)

    # Cột phải: Diagram 3 tầng kiến trúc trực quan
    right_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.75), Inches(5.683), Inches(5.1))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = WHITE
    right_box.line.color.rgb = BORDER_COLOR
    right_box.line.width = Pt(1.5)

    # Tầng 1: Frontend Client
    b_fe = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), Inches(2.0), Inches(5.083), Inches(0.95))
    b_fe.fill.solid()
    b_fe.fill.fore_color.rgb = PRIMARY_LIGHT
    b_fe.line.color.rgb = PRIMARY
    b_fe.line.width = Pt(1.5)
    tf = b_fe.text_frame
    p = tf.paragraphs[0]
    p.text = "PRESENTATION LAYER (CLIENT WEB)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "HTML5 • CSS3 Modern • JavaScript ES6+ • Bootstrap 5"
    p.font.name = FONT_BODY
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.alignment = PP_ALIGN.CENTER

    arrow1 = s5.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.45), Inches(3.05), Inches(0.45), Inches(0.32))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = ACCENT_TEAL
    arrow1.line.fill.background()

    # Tầng 2: Backend Spring Boot
    b_be = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), Inches(3.47), Inches(5.083), Inches(1.2))
    b_be.fill.solid()
    b_be.fill.fore_color.rgb = PRIMARY
    b_be.line.fill.background()
    tf = b_be.text_frame
    p = tf.paragraphs[0]
    p.text = "BUSINESS LAYER (SPRING BOOT 3 REST API)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "Spring Security 6 (JWT) • Spring Data JPA • 9 Modules Core\nApiResponse<T> • GlobalExceptionHandler • Pessimistic Lock"
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = PRIMARY_LIGHT
    p.alignment = PP_ALIGN.CENTER

    arrow2 = s5.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.45), Inches(4.77), Inches(0.45), Inches(0.32))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = ACCENT_TEAL
    arrow2.line.fill.background()

    # Tầng 3: Data & Services
    services = [
        ("POSTGRESQL 16", "Docker (Port 5432)\n15 Tables • 3NF"),
        ("CLOUDINARY CDN", "Lưu trữ hình ảnh xe\nBiên bản ODO 6 góc"),
        ("GEMINI AI FLASH", "Tư vấn chọn xe 24/7\nLưu vết chat_memory")
    ]
    for idx, (stitle, sdesc) in enumerate(services):
        c_left = Inches(7.15 + idx * 1.73)
        b_srv = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(5.18), Inches(1.63), Inches(1.35))
        b_srv.fill.solid()
        b_srv.fill.fore_color.rgb = BG_LIGHT
        b_srv.line.color.rgb = BORDER_COLOR
        b_srv.line.width = Pt(1.0)
        tf = b_srv.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.alignment = PP_ALIGN.CENTER
        p = tf.add_paragraph()
        p.text = sdesc
        p.font.name = FONT_BODY
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED
        p.alignment = PP_ALIGN.CENTER

    set_notes(s5, """LỜI THOẠI THUYẾT TRÌNH (03:00 - 03:45):
Về kiến trúc, hệ thống tách biệt Client và Server qua RESTful API chuẩn ApiResponse<T>.
Backend dùng Spring Boot 3.3.4 phân chia 9 module nghiệp vụ rõ ràng.
CSDL PostgreSQL 16 đóng gói trong Docker container đảm bảo tính nhất quán.
Hình ảnh và video được lưu trữ tối ưu trên Cloudinary CDN kèm bộ nhớ lưu vết chat AI.""")

    # =========================================================================
    # SLIDE 6: CHƯƠNG 2 - BIỂU ĐỒ USE CASE HỆ THỐNG (CHUẨN BÀI MẪU ITC)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_centered_diagram_slide(
        s6, "2.2", "Thiết Kế Chức Năng Hệ Thống (Biểu Đồ Use Case)",
        "> Biểu đồ Use Case đặt xe & vận hành thể hiện mối liên kết giữa Khách thuê, Chủ xe và Cổng VietQR:",
        "docs/image_diagram/usecase_booking_operation.png",
        "Hình 2.2 Biểu đồ Use Case quy trình đặt xe và vận hành (Booking & Operation) — Hệ thống DriveShare"
    )
    set_notes(s6, """LỜI THOẠI THUYẾT TRÌNH (03:45 - 04:30):
Đây là Sơ đồ Use Case quy trình đặt xe và vận hành cốt lõi:
- Khách thuê tạo yêu cầu, khi được duyệt sẽ thực hiện thanh toán cọc 30% qua VietQR để kích hoạt Giữ cọc Escrow & Hard Lock.
- Chủ xe khi duyệt đơn sẽ kích hoạt Khóa tạm lịch xe Soft Lock trong 45 phút.
- Hai bên cùng thực hiện Check-in giao xe có 6 góc ảnh và Check-out trả xe có tính phụ phí minh bạch.""")

    # =========================================================================
    # SLIDE 7: CHƯƠNG 2 - THIẾT KẾ CƠ SỞ DỮ LIỆU (4 BẢNG CỐT LÕI - CHUẨN MẪU ITC)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_centered_diagram_slide(
        s7, "2.3", "Thiết Kế Cơ Sở Dữ Liệu Phía Server",
        "> Xây dựng cơ sở dữ liệu sử dụng PostgreSQL 16 (15 bảng chuẩn 3NF):",
        "docs/diagrams/core_database_tables.png",
        "Hình 2.3 Cấu trúc các bảng dữ liệu hạt nhân trong hệ thống CSDL DriveShare (users, cars, rentals, payments)"
    )
    set_notes(s7, """LỜI THOẠI THUYẾT TRÌNH (04:30 - 05:20):
Về cơ sở dữ liệu, nhóm xây dựng trên PostgreSQL 16 gồm 15 bảng quan hệ chuẩn 3NF:
- Bảng users: Quản lý người dùng, phân quyền RBAC và token_version.
- Bảng cars: Quản lý thông tin xe, biển số UNIQUE và đơn giá thuê.
- Bảng rentals: Bảng hạt nhân lưu trữ toàn bộ tiến trình đơn thuê và thời hạn cọc 45 phút.
- Bảng payments: Lưu trữ mã giao dịch DSPAY, mã VietQR và trạng thái ký quỹ cọc 30%.""")

    # =========================================================================
    # SLIDE 8: CHƯƠNG 2 - VÒNG ĐỜI ĐƠN THUÊ & STATE MACHINE (CHUẨN BÀI MẪU ITC)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_centered_diagram_slide(
        s8, "2.4", "Vòng Đời Đơn Thuê & Ký Quỹ VietQR (State Machine)",
        "> Chu trình điều khiển trạng thái tự động loại bỏ lỗi Overbooking và bảo vệ dòng tiền cọc:",
        "docs/image_diagram/VietQR Rental Booking Lifecycle-2026-10-28-150558.png",
        "Hình 2.4 Máy trạng thái điều khiển vòng đời đơn thuê và ký quỹ cọc VietQR 30% — DriveShare"
    )
    set_notes(s8, """LỜI THOẠI THUYẾT TRÌNH (05:20 - 06:10):
Toàn bộ quy trình đặt xe được điều khiển bởi State Machine chuyên sâu:
- Đơn tạo mới PENDING_APPROVAL: Chủ xe có 60 phút để duyệt, quá giờ tự hủy EXPIRED.
- Duyệt đơn kích hoạt WAITING_PAYMENT (Soft Lock 45 phút), các đơn trùng lịch tự sang ON_HOLD.
- Cọc 30% VietQR thành công chuyển sang CONFIRMED (Hard Lock) và tự động từ chối các đơn ON_HOLD.
- Nếu khách hết 45 phút không cọc, đơn chuyển EXPIRED và hệ thống tự động phục hồi đơn ON_HOLD về Chờ duyệt!""")

    # =========================================================================
    # SLIDE 9: CHƯƠNG 2 - ĐIỂM SÁNG 01: KHÓA BI QUAN PESSIMISTIC_WRITE (CODE CARD)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_slide_header(s9, "2.5", "Điểm Sáng Kỹ Thuật: Khóa Bi Quan PESSIMISTIC_WRITE",
                     "> Cơ chế chống xung đột đồng thời (Race Condition) triệt để ở tầng CSDL:")

    # Cột trái: Card giải thích
    left_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.8), Inches(5.1))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = BORDER_COLOR
    left_card.line.width = Pt(1.5)

    tf_l = left_card.text_frame
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.35)
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "GIẢI PHÁP CHỐNG XUNG ĐỘT ĐỒNG THỜI"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    bullets_s9 = [
        "Vấn đề Race Condition: Nhiều luồng cùng duyệt giữ xe hoặc cùng hoàn cọc gây lỗi trùng lặp.",
        "Giải pháp trong code: Áp dụng @Lock(LockModeType.PESSIMISTIC_WRITE) tại CarRepository.",
        "Cơ chế SQL: Phát sinh lệnh SELECT ... FOR UPDATE khóa độc quyền bản ghi ở tầng CSDL.",
        "An toàn tuyệt đối: Giữ khóa cho đến khi commit transaction; loại bỏ 100% nguy cơ Overbooking.",
        "Hoàn cọc Idempotent: Khóa bản ghi khiếu nại tại RentalDisputeRepository chống hoàn tiền 2 lần."
    ]
    for b in bullets_s9:
        p = tf_l.add_paragraph()
        p.text = "•  " + b
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(9)

    # Cột phải: Code Box Dark Theme
    code_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.75), Inches(5.683), Inches(5.1))
    code_box.fill.solid()
    code_box.fill.fore_color.rgb = CODE_BG
    code_box.line.color.rgb = PRIMARY
    code_box.line.width = Pt(2.0)

    tf_c = code_box.text_frame
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = Inches(0.35)
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "// CarRepository.java (Spring Data JPA)"
    p.font.name = "Consolas"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(4)

    p = tf_c.add_paragraph()
    p.text = "@Lock(LockModeType.PESSIMISTIC_WRITE)"
    p.font.name = "Consolas"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CODE_YELLOW
    p.space_after = Pt(2)

    p = tf_c.add_paragraph()
    p.text = '@Query("SELECT c FROM Car c WHERE c.id = :id")'
    p.font.name = "Consolas"
    p.font.size = Pt(12)
    p.font.color.rgb = CODE_GREEN
    p.space_after = Pt(2)

    p = tf_c.add_paragraph()
    p.text = "Optional<Car> findCarForUpdate(@Param(\"id\") Long id);"
    p.font.name = "Consolas"
    p.font.size = Pt(12.5)
    p.font.color.rgb = WHITE
    p.space_after = Pt(16)

    p = tf_c.add_paragraph()
    p.text = "-- SQL Generated by Hibernate (PostgreSQL 16):"
    p.font.name = "Consolas"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(4)

    p = tf_c.add_paragraph()
    p.text = "SELECT c1_0.id, c1_0.license_plate, c1_0.price\nFROM cars c1_0\nWHERE c1_0.id = ?\nFOR UPDATE;"
    p.font.name = "Consolas"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CODE_YELLOW
    p.space_after = Pt(16)

    p = tf_c.add_paragraph()
    p.text = "=> Khóa độc quyền hàng dữ liệu ngay lập tức.\n=> Luồng thứ 2 phải xếp hàng chờ cho đến khi commit.\n=> Triệt tiêu 100% rủi ro Overbooking!"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = CODE_GREEN

    set_notes(s9, """LỜI THOẠI THUYẾT TRÌNH (06:10 - 07:10):
Đây là Điểm sáng kỹ thuật then chốt: Xử lý Race Condition bằng Khóa bi quan.
Lệnh SELECT FOR UPDATE khóa trực tiếp hàng dữ liệu xe trong PostgreSQL trong suốt transaction duyệt đơn, đảm bảo không thể có 2 đơn cùng được chốt một lúc.
Khi Admin xử lý khiếu nại, bản ghi tranh chấp cũng được khóa để chống hoàn tiền 2 lần.""")

    # =========================================================================
    # SLIDE 10: CHƯƠNG 2 - ĐIỂM SÁNG 02: KHÓA 2 TẦNG (TWO-TIER LOCKING - SEQUENCE)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_centered_diagram_slide(
        s10, "2.6", "Điểm Sáng Kỹ Thuật: Cơ Chế Khóa 2 Tầng (Two-Tier Locking)",
        "> Sơ đồ tuần tự quy trình đặt xe, ký quỹ cọc 30% VietQR và tự động phục hồi đơn chờ:",
        "docs/diagrams/04_sequence_booking_and_deposit.png",
        "Hình 2.6 Sơ đồ tuần tự quy trình đặt xe và khóa 2 tầng (Soft Lock 45m & Hard Lock Ký Quỹ) — DriveShare"
    )
    set_notes(s10, """LỜI THOẠI THUYẾT TRÌNH (07:10 - 08:00):
Tiếp nối là Cơ chế Khóa 2 tầng linh hoạt:
Soft Lock 45 phút cho khách thời gian chuyển khoản cọc 30% mà không sợ mất chỗ.
Khi có cọc, Hard Lock kích hoạt chốt xe vĩnh viễn.
Nếu quá 45 phút khách không cọc, Cronjob tự động giải phóng xe và phục hồi các đơn ON_HOLD về Chờ duyệt để đón khách tiếp theo.""")

    # =========================================================================
    # SLIDE 11: CHƯƠNG 2 - ĐIỂM SÁNG 03: BẢO MẬT JWT & AI HYBRID CHỊU LỖI
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_slide_header(s11, "2.7", "Điểm Sáng Kỹ Thuật: Bảo Mật JWT & AI Hybrid Chịu Lỗi",
                     "> Cơ chế bảo mật phiên Stateless và phân hệ AI hoạt động tin cậy cả Online lẫn Offline:")

    left_card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.8), Inches(5.1))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = BORDER_COLOR
    left_card.line.width = Pt(1.5)

    tf_l = left_card.text_frame
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.35)
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "AN TOÀN BẢO MẬT & FAULT-TOLERANCE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    bullets_s11 = [
        "Stateless JWT: Dùng jjwt 0.12.6, mã hóa mật khẩu BCrypt, phân quyền phương thức RBAC.",
        "Token Versioning: Trường token_version giúp thu hồi token tức thì ngay khi đổi mật khẩu.",
        "Chống Brute-force: Tự động khóa tài khoản locked_until sau 5 lần nhập sai mật khẩu liên tiếp.",
        "Google Gemini 2.5 Flash: Kết nối REST API trực tiếp qua Java HttpClient kèm ngữ cảnh CSDL.",
        "Heuristic Regex Fallback: Tự động bóc tách ngày, giá, số chỗ khi mất mạng hoặc chưa có API Key."
    ]
    for b in bullets_s11:
        p = tf_l.add_paragraph()
        p.text = "•  " + b
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(9)

    # Cột phải: Ảnh chụp AI Chatbot
    box_r = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.75), Inches(5.683), Inches(5.1))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = WHITE
    box_r.line.color.rgb = BORDER_COLOR
    box_r.line.width = Pt(1.5)

    img_ai = "docs/screenshots/02_ai_chatbot_dialog.png"
    if os.path.exists(img_ai):
        try:
            with Image.open(img_ai) as img:
                orig_w, orig_h = img.size
            aspect = orig_w / orig_h
            dest_w = Inches(5.3)
            dest_h = dest_w / aspect
            dest_l = Inches(7.04)
            dest_t = Inches(1.75) + (Inches(5.1) - dest_h) / 2
            s11.shapes.add_picture(img_ai, dest_l, dest_t, dest_w, dest_h)
        except Exception as e:
            print(f"Lỗi ảnh AI: {e}")

    set_notes(s11, """LỜI THOẠI THUYẾT TRÌNH (08:00 - 08:30):
Về bảo mật, cơ chế token_version cho phép thu hồi JWT ngay khi đổi mật khẩu, tự khóa tài khoản sau 5 lần đăng nhập sai.
Phân hệ AI hoạt động theo cơ chế chịu lỗi Fault-Tolerant: Bình thường gọi Gemini 2.5 Flash; khi mất mạng bộ máy Heuristic Regex tự kích hoạt để tư vấn xe và tính cọc 30% mượt mà.""")

    # =========================================================================
    # SLIDE 12: CHƯƠNG 3 - PHÂN HỆ KHÁCH THUÊ (3 ẢNH MOBILE THEO YÊU CẦU)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    renter_screens = [
        ("docs/screenshots/mobile_renter_search.png",
         "1. Tìm Xe & Bộ Lọc",
         "Lọc theo Tỉnh, ngày thuê, khoảng giá và số chỗ"),
        ("docs/screenshots/mobile_renter_vietqr.png",
         "2. Ký Quỹ Cọc 30% VietQR",
         "Tự sinh mã QR ngân hàng & đếm ngược Soft Lock 45m"),
        ("docs/screenshots/07_mobile_375px_dispute_modal.png",
         "3. Khiếu Nại Sự Cố (Dispute)",
         "Gửi video bằng chứng tại chỗ để hoàn 100% cọc")
    ]
    add_three_mobile_screens_slide(
        s12, "3.1", "Kết Quả Thực Nghiệm: Phân Hệ Khách Thuê (Renter)",
        "> Giao diện Mobile Responsive trực quan, tối ưu hành trình trải nghiệm người dùng:",
        renter_screens
    )
    set_notes(s12, """LỜI THOẠI THUYẾT TRÌNH (08:30 - 09:30):
Bước sang Chương 3 thực nghiệm, đây là 3 màn hình cốt lõi của Khách thuê trên Mobile:
1. Màn hình tìm xe với bộ lọc đa tiêu chí trực quan.
2. Màn hình thanh toán ký quỹ 30% VietQR tự động sinh mã kèm đồng hồ đếm ngược Soft Lock 45 phút.
3. Màn hình gửi khiếu nại sự cố kèm video bằng chứng để nhận lại 100% tiền cọc nếu chủ xe không giao xe.""")

    # =========================================================================
    # SLIDE 13: CHƯƠNG 3 - PHÂN HỆ CHỦ XE (3 ẢNH GIAO DIỆN THEO YÊU CẦU)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    owner_screens = [
        ("docs/screenshots/03_owner_cars_dashboard.png",
         "1. Quản Lý Đội Xe",
         "Bật/Tắt trạng thái hoạt động & cập nhật giá thuê"),
        ("docs/screenshots/proof_03_vi_two_way_inspection_modal_ui.png",
         "2. Biên Bản Bàn Giao 2 Chiều",
         "Check-in/out ghi nhận ODO, vạch xăng & 6 góc ảnh"),
        ("docs/screenshots/04_owner_earnings_wallet.png",
         "3. Ví Doanh Thu Chủ Xe",
         "Thống kê số dư khả dụng, nhận 90% doanh thu cuốc xe")
    ]
    add_three_mobile_screens_slide(
        s13, "3.2", "Kết Quả Thực Nghiệm: Phân Hệ Chủ Xe (Car Owner)",
        "> Bộ công cụ số hóa toàn diện giúp chủ xe quản lý tài sản và tối ưu doanh thu:",
        owner_screens
    )
    set_notes(s13, """LỜI THOẠI THUYẾT TRÌNH (09:30 - 10:30):
Đối với phân hệ Chủ xe, nhóm triển khai 3 chức năng thực nghiệm then chốt:
1. Dashboard quản lý đội xe với nút bật/tắt nhận khách tức thì.
2. Biên bản bàn giao điện tử 2 chiều: Nhập số km ODO, vạch xăng và chụp 6 góc ảnh làm bằng chứng đối soát.
3. Ví doanh thu minh bạch: Tự động quyết toán 90% tiền thuê cho chủ xe sau khi cuốc xe hoàn tất.""")

    # =========================================================================
    # SLIDE 14: CHƯƠNG 3 - PHÂN HỆ QUẢN TRỊ (3 ẢNH GIAO DIỆN ADMIN THEO YÊU CẦU)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    admin_screens = [
        ("docs/screenshots/05_admin_portal_overview.png",
         "1. Dashboard Tổng Quan Sàn",
         "Theo dõi thời gian thực số xe, thành viên & dòng tiền cọc"),
        ("docs/screenshots/05_admin_dual_cards_cccd_and_bank.png",
         "2. Thẩm Định eKYC Giấy Tờ",
         "Đối soát ảnh chụp CCCD, GPLX và tài khoản ngân hàng"),
        ("docs/screenshots/02_dispute_modal_open.png",
         "3. Dispute & Chế Tài Điều 175",
         "Phân xử hoàn 100% cọc & xuất Hồ sơ vi phạm 1-click")
    ]
    add_three_mobile_screens_slide(
        s14, "3.3", "Kết Quả Thực Nghiệm: Phân Hệ Quản Trị (Admin Portal)",
        "> Trung tâm điều hành kiểm soát rủi ro, phân xử khiếu nại và bảo đảm an ninh pháp lý:",
        admin_screens
    )
    set_notes(s14, """LỜI THOẠI THUYẾT TRÌNH (10:30 - 11:30):
Cổng Quản trị Admin đảm nhiệm 3 vai trò sống còn:
1. Dashboard giám sát toàn bộ hoạt động và doanh thu ký quỹ sàn.
2. Thẩm định trực quan hồ sơ eKYC CCCD và GPLX của khách thuê trước khi cho phép đặt xe.
3. Trung tâm phân xử Dispute giải ngân hoàn cọc 100%, phạt chủ xe vi phạm 500k và tự động khóa nick, xuất hồ sơ theo Điều 175 BLHS nếu quá hạn 24h.""")

    # =========================================================================
    # SLIDE 15: CHƯƠNG 3 - QUY TRÌNH BÀN GIAO XE 2 CHIỀU (SEQUENCE DIAGRAM)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_centered_diagram_slide(
        s15, "3.4", "Quy Trình Nghiệm Thu Xe 2 Chiều & Hợp Đồng Điện Tử",
        "> Sơ đồ tuần tự nghiệp vụ Check-in & Check-out triệt tiêu hoàn toàn tranh chấp hư hỏng:",
        "docs/diagrams/05_sequence_two_way_handover.png",
        "Hình 3.4 Sơ đồ tuần tự quy trình bàn giao xe hai chiều (Check-in & Check-out) — DriveShare"
    )
    set_notes(s15, """LỜI THOẠI THUYẾT TRÌNH (11:30 - 12:15):
Điểm đột phá tiếp theo là Quy trình nghiệm thu xe 2 chiều:
Cả lúc giao xe và nhận xe đều bắt buộc nhập ODO, vạch xăng và chụp 6 góc ảnh.
Dữ liệu lưu vào bảng rental_inspections tạo bằng chứng đối soát minh bạch, loại bỏ hoàn toàn các vụ cãi vã đền bù hư hỏng.""")

    # =========================================================================
    # SLIDE 16: CHƯƠNG 3 - KIỂM THỬ TỰ ĐỘNG 188 TESTS (IMAGE)
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_background(s16)
    add_slide_header(s16, "3.5", "Chiến Lược Kiểm Thử Tự Động Toàn Diện (188 Tests)",
                     "> Đảm bảo chất lượng phần mềm và độ chịu tải đồng thời trước khi bàn giao:")

    left_card = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.8), Inches(5.1))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = BORDER_COLOR
    left_card.line.width = Pt(1.5)

    tf_l = left_card.text_frame
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.35)
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "CHẤT LƯỢNG PHẦN MỀM"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    bullets_s16 = [
        "188 bài test tự động: Đạt tỷ lệ PASS 100% (Unit Test, Integration Test, Concurrency Test).",
        "Unit & Service Tests: Kiểm thử công thức phạt trễ hạn lũy tiến, tính cọc 30%, kiểm tra trùng lịch.",
        "MockMvc Integration Tests: Kiểm thử toàn bộ các luồng xác thực JWT và phân quyền API.",
        "Concurrency Tests: Giả lập đa luồng đặt xe cùng lúc, xác nhận loại bỏ hoàn toàn Overbooking.",
        "Kiểm soát chất lượng: Tích hợp JaCoCo đo độ bao phủ mã nguồn trước khi merge Pull Request."
    ]
    for b in bullets_s16:
        p = tf_l.add_paragraph()
        p.text = "•  " + b
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(9)

    box_r = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.75), Inches(5.683), Inches(5.1))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = WHITE
    box_r.line.color.rgb = BORDER_COLOR
    box_r.line.width = Pt(1.5)

    img_qa = "docs/screenshots/08_qa_test_suite_matrix.png"
    if os.path.exists(img_qa):
        try:
            with Image.open(img_qa) as img:
                orig_w, orig_h = img.size
            aspect = orig_w / orig_h
            dest_h = Inches(4.8)
            dest_w = dest_h * aspect
            if dest_w > Inches(5.4):
                dest_w = Inches(5.4)
                dest_h = dest_w / aspect
            dest_l = Inches(6.85) + (Inches(5.683) - dest_w) / 2
            dest_t = Inches(1.75) + (Inches(5.1) - dest_h) / 2
            s16.shapes.add_picture(img_qa, dest_l, dest_t, dest_w, dest_h)
        except Exception as e:
            print(f"Lỗi ảnh QA: {e}")

    set_notes(s16, """LỜI THOẠI THUYẾT TRÌNH (12:15 - 13:00):
Để đảm bảo chất lượng, nhóm xây dựng 188 bài test tự động:
Từ Unit Test logic tính tiền, phạt trễ hạn;
MockMvc kiểm thử phân quyền;
Đến Concurrency Test giả lập đa luồng đặt xe cùng lúc để chứng minh tính an toàn của Khóa bi quan.
Toàn bộ 188 tests đều pass 100%.""")

    # =========================================================================
    # SLIDE 17: KẾT LUẬN - TỔNG KẾT KẾT QUẢ ĐẠT ĐƯỢC (3 CARDS)
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    card_data_s17 = [
        ("KẾT QUẢ 01", "Hoàn Thành 5 Sprints Scrum",
         [
             "Tuân thủ chặt chẽ mô hình Agile Scrum từ Sprint 1 đến Sprint 5.",
             "Hoàn thành 100% User Stories cam kết theo đề bài thực tập K25.",
             "Thực hiện họp Daily Scrum và Retrospective đúng định kỳ."
         ]),
        ("KẾT QUẢ 02", "Kiến Trúc Thực Tế Chuẩn Mực",
         [
             "Áp dụng thành công Khóa bi quan CSDL, Two-Tier Lock và VietQR 30%.",
             "Biên bản bàn giao điện tử 2 chiều triệt tiêu tranh chấp hư hỏng.",
             "Trợ lý AI hoạt động ổn định và sẵn sàng cả online lẫn offline."
         ]),
        ("KẾT QUẢ 03", "Quy Trình Doanh Nghiệp",
         [
             "Mô hình Git Feature-branch và Pull Request Review nghiêm ngặt.",
             "188 automated tests bảo vệ chất lượng mã nguồn trước khi release.",
             "Hệ thống vận hành trơn tru và sẵn sàng demo trực tiếp."
         ])
    ]
    add_three_cards_slide(s17, "KẾT LUẬN", "Tổng Kết Kết Quả Đạt Được Của Đề Tài",
                          "> Những thành quả nổi bật sau quá trình nghiên cứu và triển khai dự án:",
                          card_data_s17)
    set_notes(s17, """LỜI THOẠI THUYẾT TRÌNH (13:00 - 13:45):
Tóm lại, sau 5 Sprints Agile Scrum, nhóm chúng em đã đạt được các kết quả trọng tâm:
1. Xây dựng nền tảng thuê xe P2P an toàn, khép kín.
2. Áp dụng thành công các kỹ thuật chuyên sâu: Khóa bi quan CSDL, ký quỹ VietQR 30%, bàn giao 2 chiều và AI tư vấn.
3. Tuân thủ quy trình phát triển doanh nghiệp chuẩn chỉ với 188 bài test tự động.""")

    # =========================================================================
    # SLIDE 18: KẾT LUẬN - HẠN CHẾ & HƯỚNG PHÁT TRIỂN (LƯỚI 2X2 ĐƠN GIẢN)
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    card_data_s18 = [
        ("HẠN CHẾ 01", "Xác Nhận VietQR",
         "Hiện tại: Xác nhận cọc qua nút bấm mô phỏng Webhook trên màn hình.\nHướng tới: Tích hợp Webhook biến động số dư ngân hàng tự động (Casso / SePay)."),
        ("HẠN CHẾ 02", "Thẩm Định eKYC",
         "Hiện tại: Admin duyệt bằng mắt trên ảnh chụp CCCD và bằng lái xe.\nHướng tới: Tích hợp AI OCR quét thông tin giấy tờ và so khớp khuôn mặt sinh trắc."),
        ("HẠN CHẾ 03", "Thông Báo Realtime",
         "Hiện tại: Sử dụng cơ chế REST Polling định kỳ để cập nhật chuông.\nHướng tới: Nâng cấp lên kết nối hai chiều WebSocket STOMP Broker."),
        ("HẠN CHẾ 04", "Bản Đồ Tương Tác",
         "Hiện tại: Quản lý tọa độ trong DB, giao diện lọc theo Tỉnh/Thành.\nHướng tới: Nhúng bản đồ Mapbox/Leaflet tìm xe theo bán kính 5 - 10km.")
    ]
    add_four_cards_slide(s18, "KẾT LUẬN", "Những Hạn Chế Còn Tồn Tại & Hướng Phát Triển",
                         "> Đánh giá trung thực hiện trạng và định hướng nâng cấp hệ thống:",
                         card_data_s18)
    set_notes(s18, """LỜI THOẠI THUYẾT TRÌNH (13:45 - 14:35):
Với tinh thần cầu thị và trung thực khoa học, nhóm xin chủ động báo cáo 4 hạn chế hiện tại:
1. Xác nhận chuyển khoản VietQR đang dùng nút mô phỏng -> Hướng tới Webhook Casso.
2. eKYC đang do Admin duyệt bằng mắt -> Hướng tới AI OCR nhận diện CCCD.
3. Thông báo dùng REST Polling -> Hướng tới WebSocket STOMP.
4. Danh sách xe lọc theo Tỉnh/Thành -> Hướng tới nhúng bản đồ tương tác Mapbox.""")

    # =========================================================================
    # SLIDE 19: LỜI CẢM ƠN & PHIÊN VẤN ĐÁP Q&A (SPLIT)
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    set_slide_background(s19)
    add_slide_header(s19, "Q&A", "Phiên Vấn Đáp & Lắng Nghe Góp Ý Của Hội Đồng",
                     "> Lời tri ân chân thành và phân công thành viên giải đáp phản biện:")

    card_ty = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.8), Inches(5.1))
    card_ty.fill.solid()
    card_ty.fill.fore_color.rgb = WHITE
    card_ty.line.color.rgb = BORDER_COLOR
    card_ty.line.width = Pt(1.5)

    tf = card_ty.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.35)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "LỜI TRI ÂN SÂU SẮC"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "Nhóm xin trân trọng cảm ơn quý Thầy, Cô trong Hội đồng phản biện đã dành thời gian quý báu theo dõi bài báo cáo của nhóm.\n\nĐặc biệt, nhóm xin gửi lời tri ân sâu sắc nhất đến Thầy Trương Thanh Bình đã luôn đồng hành, định hướng và chỉ bảo tận tình cho nhóm trong suốt quá trình hoàn thiện đề án tốt nghiệp này."
    p.font.name = FONT_BODY
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_MAIN

    card_qa = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.75), Inches(5.683), Inches(5.1))
    card_qa.fill.solid()
    card_qa.fill.fore_color.rgb = WHITE
    card_qa.line.color.rgb = BORDER_COLOR
    card_qa.line.width = Pt(1.5)

    tf = card_qa.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.35)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "PHIÊN VẤN ĐÁP & HỎI ĐÁP (Q&A)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "Nhóm sẵn sàng lắng nghe những góp ý quý báu và giải đáp mọi câu hỏi phản biện từ quý Thầy, Cô!\n\nĐại diện trả lời: Nguyễn Duy Bảo (Tech Lead)\nCùng sự phối hợp của các thành viên chuyên trách:\n•  Nguyễn Chí Tín: Thanh toán VietQR & Giao diện Web\n•  Lâm Chí Vĩ: Backend Core & Khóa bi quan CSDL\n•  Đỗ Ngọc Phát: Trợ lý AI Gemini & Bảo mật JWT\n•  Nguyễn Duy Quân: Kiểm thử 188 Tests & Xử phạt quá hạn\n•  NguyễnTấn Lộc Khiêm: Thiết kế Kiến trúc & Use Case"
    p.font.name = FONT_BODY
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_MAIN

    set_notes(s19, """LỜI THOẠI THUYẾT TRÌNH (14:35 - 15:00):
Bài thuyết trình báo cáo đề tài của nhóm chúng em đến đây là kết thúc.
Nhóm xin gửi lời cảm ơn chân thành nhất đến quý Thầy Cô trong Hội đồng và đặc biệt là Thầy Trương Thanh Bình.
Sau đây, nhóm chúng em xin được lắng nghe những ý kiến đóng góp quý báu và sẵn sàng trả lời các câu hỏi phản biện từ quý Thầy Cô. Em xin cảm ơn!""")

    # Lưu đồng thời ra cả 2 file
    paths_to_save = [
        "docs/Slide_Bao_Ve_Do_An_DriveShare_K4_Moi_Nhat.pptx",
        "docs/Slide_Bao_Ve_Do_An_DriveShare_K4.pptx"
    ]
    for pth in paths_to_save:
        try:
            prs.save(pth)
            print(f"[OK] Xuất thành công {len(prs.slides)} slides chuẩn ITC K25 vào: {pth}")
        except Exception as e:
            print(f"[!] Không thể lưu vào {pth}: {e}")


if __name__ == "__main__":
    build_presentation()
