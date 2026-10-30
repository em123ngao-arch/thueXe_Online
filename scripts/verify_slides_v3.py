import sys
from pptx import Presentation

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation('docs/Slide_Bao_Ve_Do_An_DriveShare_K4_Moi_Nhat.pptx')
print(f"=== KIỂM TRA 19 SLIDES (Tổng: {len(prs.slides)} slides) ===")

for i, slide in enumerate(prs.slides, 1):
    texts = [p.text.strip() for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text.strip()]
    title = texts[0] if texts else "EMPTY"
    img_shapes = [s for s in slide.shapes if s.shape_type == 13]
    has_notes = slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip())
    
    font_sizes = []
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.font.size:
                    font_sizes.append(p.font.size.pt)
    min_sz = min(font_sizes) if font_sizes else 0
    max_sz = max(font_sizes) if font_sizes else 0

    print(f"Slide {i:02d} | Shapes: {len(slide.shapes):2d} | Imgs: {len(img_shapes)} | Font: [{min_sz:4.1f} - {max_sz:4.1f}] | Notes: {str(has_notes):5s} | {title[:50]}")

